#!/usr/bin/env python
"""Batch RAG answers for question/Demo问题库.xlsx → Demo问题库_回答结果.xlsx + .jsonl."""

from __future__ import annotations

import argparse
import asyncio
import json
import math
import os
import sys
from functools import partial
from pathlib import Path

from dotenv import load_dotenv

# Project root
_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(_ROOT))
load_dotenv(dotenv_path=_ROOT / ".env", override=False)

from lightrag import LightRAG
from lightrag.llm.openai import openai_complete_if_cache, openai_embed
from lightrag.utils import EmbeddingFunc, logger
from raganything import RAGAnything, RAGAnythingConfig
from raganything.local_hf_embedding import (
    ensure_hf_home_from_repo_fallback,
    make_local_hf_embedding_func,
)


def _json_safe(val):
    if isinstance(val, float) and (math.isnan(val) or math.isinf(val)):
        return None
    return val


def _resolve_keys() -> tuple[str, str]:
    llm_key = (
        os.getenv("OPENAI_API_KEY", "").strip()
        or os.getenv("LLM_BINDING_API_KEY", "").strip()
    )
    emb_key = os.getenv("EMBEDDING_API_KEY", "").strip() or llm_key
    return llm_key, emb_key


async def run_batch(
    input_xlsx: Path,
    out_xlsx: Path,
    out_jsonl: Path,
    working_dir: Path,
    mode: str,
    delay_seconds: float,
    embedding_func_max_async: int,
    embedding_batch_num: int,
) -> None:
    import openpyxl

    ensure_hf_home_from_repo_fallback(_ROOT)

    llm_key, emb_key = _resolve_keys()
    if not llm_key:
        raise SystemExit(
            "Set OPENAI_API_KEY or LLM_BINDING_API_KEY in the environment or .env"
        )

    embedding_backend = os.getenv("EMBEDDING_BACKEND", "openai").strip().lower()

    base_url = os.getenv("LLM_BINDING_HOST", "").strip() or None
    emb_host = os.getenv("EMBEDDING_BINDING_HOST", "").strip()
    embedding_base_url = emb_host if emb_host else base_url
    if embedding_backend != "hf":
        if emb_host and not os.getenv("EMBEDDING_API_KEY", "").strip():
            raise SystemExit(
                "EMBEDDING_BINDING_HOST is set; set EMBEDDING_API_KEY to the key for that host "
                "(e.g. OpenAI platform key when using text-embedding-3-small on api.openai.com)."
            )
    llm_model = os.getenv("LLM_MODEL", "gpt-4o-mini")
    vision_model = os.getenv("VISION_MODEL", llm_model)
    if embedding_backend == "hf":
        embedding_dim = int(os.getenv("EMBEDDING_DIM", "1024"))
        embedding_model = os.getenv("EMBEDDING_MODEL", "BAAI/bge-m3")
    else:
        embedding_dim = int(os.getenv("EMBEDDING_DIM", "1536"))
        embedding_model = os.getenv(
            "EMBEDDING_MODEL", "text-embedding-3-small"
        ).strip()

    config = RAGAnythingConfig(
        working_dir=str(working_dir),
        parser=os.getenv("PARSER", "mineru"),
        parse_method="auto",
        enable_image_processing=True,
        enable_table_processing=True,
        enable_equation_processing=True,
    )

    def llm_model_func(prompt, system_prompt=None, history_messages=[], **kwargs):
        return openai_complete_if_cache(
            llm_model,
            prompt,
            system_prompt=system_prompt,
            history_messages=history_messages,
            api_key=llm_key,
            base_url=base_url,
            **kwargs,
        )

    def vision_model_func(
        prompt,
        system_prompt=None,
        history_messages=None,
        image_data=None,
        messages=None,
        **kwargs,
    ):
        if history_messages is None:
            history_messages = []
        if messages:
            return openai_complete_if_cache(
                vision_model,
                "",
                system_prompt=None,
                history_messages=[],
                messages=messages,
                api_key=llm_key,
                base_url=base_url,
                **kwargs,
            )
        if image_data:
            return openai_complete_if_cache(
                vision_model,
                "",
                system_prompt=None,
                history_messages=[],
                messages=[
                    {"role": "system", "content": system_prompt}
                    if system_prompt
                    else None,
                    {
                        "role": "user",
                        "content": [
                            {"type": "text", "text": prompt},
                            {
                                "type": "image_url",
                                "image_url": {
                                    "url": f"data:image/jpeg;base64,{image_data}",
                                },
                            },
                        ],
                    },
                ],
                api_key=llm_key,
                base_url=base_url,
                **kwargs,
            )
        return llm_model_func(prompt, system_prompt, history_messages, **kwargs)

    if embedding_backend == "hf":
        embedding_func = make_local_hf_embedding_func(
            embedding_dim,
            embedding_model=embedding_model,
        )
    else:
        embedding_func = EmbeddingFunc(
            embedding_dim=embedding_dim,
            max_token_size=8192,
            func=partial(
                openai_embed.func,
                model=embedding_model,
                api_key=emb_key,
                base_url=embedding_base_url,
            ),
        )

    lightrag = LightRAG(
        working_dir=str(working_dir),
        llm_model_func=llm_model_func,
        embedding_func=embedding_func,
        enable_llm_cache=True,
        embedding_func_max_async=embedding_func_max_async,
        embedding_batch_num=embedding_batch_num,
    )
    await lightrag.initialize_storages()

    rag = RAGAnything(
        config=config,
        lightrag=lightrag,
        llm_model_func=llm_model_func,
        vision_model_func=vision_model_func,
        embedding_func=embedding_func,
    )

    wb_in = openpyxl.load_workbook(input_xlsx)
    ws_in = wb_in.active
    headers = [c.value for c in next(ws_in.iter_rows(min_row=1, max_row=1))]
    try:
        q_col = headers.index("问题") + 1
    except ValueError as e:
        raise SystemExit(f'Column "问题" not found; got {headers!r}') from e

    if "RAG回答" not in headers:
        r_col = len(headers) + 1
        ws_in.cell(row=1, column=r_col, value="RAG回答")
    else:
        r_col = headers.index("RAG回答") + 1

    records: list[dict] = []
    rows = list(ws_in.iter_rows(min_row=2, values_only=False))
    for row_cells in rows:
        q_cell = row_cells[q_col - 1]
        question = q_cell.value
        if question is None or (isinstance(question, str) and not question.strip()):
            continue
        question = str(question).strip()
        row_num = q_cell.row
        logger.info(f"[{row_num}] Q: {question[:80]}...")
        answer = await rag.aquery(
            question,
            mode=mode,
            vlm_enhanced=False,
        )
        if answer is None:
            answer = ""
        ws_in.cell(row=row_num, column=r_col, value=answer)
        row_vals = [c.value for c in row_cells]
        rec = {
            headers[i]: _json_safe(row_vals[i])
            if i < len(row_vals)
            else None
            for i in range(len(headers))
        }
        rec["RAG回答"] = answer
        records.append(rec)
        if delay_seconds > 0:
            await asyncio.sleep(delay_seconds)

    wb_in.save(out_xlsx)
    with open(out_jsonl, "w", encoding="utf-8") as f:
        for rec in records:
            f.write(
                json.dumps(rec, ensure_ascii=False, default=str) + "\n"
            )

    logger.info(f"Wrote {len(records)} rows → {out_xlsx} and {out_jsonl}")


def main() -> None:
    p = argparse.ArgumentParser()
    p.add_argument(
        "--input",
        type=Path,
        default=_ROOT / "question" / "Demo问题库.xlsx",
    )
    p.add_argument(
        "--out-xlsx",
        type=Path,
        default=_ROOT / "question" / "Demo问题库_回答结果.xlsx",
    )
    p.add_argument(
        "--out-jsonl",
        type=Path,
        default=_ROOT / "question" / "Demo问题库_回答结果.jsonl",
    )
    p.add_argument(
        "-w",
        "--working-dir",
        type=Path,
        default=_ROOT / "rag_storage",
    )
    p.add_argument(
        "--mode",
        default=os.getenv("RAG_QUERY_MODE", "mix"),
        help="LightRAG query mode (default: mix)",
    )
    p.add_argument(
        "--delay",
        type=float,
        default=float(os.getenv("DEMO_QA_QUERY_DELAY_SECONDS", "4")),
        help="Seconds to sleep after each answer (helps avoid embedding rate limits).",
    )
    p.add_argument(
        "--embedding-max-async",
        type=int,
        default=int(os.getenv("EMBEDDING_FUNC_MAX_ASYNC", "1")),
        help="Concurrency for embedding RPCs (lower helps with Ark rate limits).",
    )
    p.add_argument(
        "--embedding-batch-num",
        type=int,
        default=int(os.getenv("EMBEDDING_BATCH_NUM", "1")),
    )
    args = p.parse_args()
    asyncio.run(
        run_batch(
            args.input,
            args.out_xlsx,
            args.out_jsonl,
            args.working_dir,
            args.mode,
            args.delay,
            args.embedding_max_async,
            args.embedding_batch_num,
        )
    )


if __name__ == "__main__":
    main()
