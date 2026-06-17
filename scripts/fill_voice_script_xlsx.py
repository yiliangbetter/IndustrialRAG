#!/usr/bin/env python3
"""Fill RAG answers from voice script test report into Excel 系统回答 column."""

from __future__ import annotations

import argparse
import re
import sys
from datetime import datetime
from pathlib import Path

_ROOT = Path(__file__).resolve().parent.parent
_DEFAULT_XLSX = Path(r"d:\dev\docs\AI智能语音话术 解析.xlsx")
_DEFAULT_REPORT = _ROOT / "logs" / "voice_script_tests" / "20260614_183740_技术.md"
_SHEET = "技术问题"
_COL_SYSTEM_ANSWER = 7  # 1-based: G = 系统回答
_HEADER_ROW = 2
_DATA_START = 3


def _parse_report(path: Path) -> dict[str, dict]:
    text = path.read_text(encoding="utf-8")
    entries: dict[str, dict] = {}
    pattern = re.compile(
        r"^## VS(\d+)\s+[✓✗]\s*\n+"
        r"\*\*标准问题\*\*：(.+?)\n"
        r"\*\*测试问法\*\*：(.+?)\n"
        r".*?"
        r"\*\*系统应答\*\*\s*\n+"
        r"(.*?)"
        r"(?=\n### References|\n## VS\d+|\Z)",
        re.MULTILINE | re.DOTALL,
    )
    for m in pattern.finditer(text):
        entries[m.group(2).strip()] = {
            "id": int(m.group(1)),
            "standard_question": m.group(2).strip(),
            "utterance": m.group(3).strip(),
            "answer": m.group(4).strip(),
        }
    return entries


def _load_workbook(path: Path):
    try:
        from openpyxl import load_workbook
    except ImportError:
        import subprocess

        subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
        from openpyxl import load_workbook
    return load_workbook(path)


def fill_xlsx(
    xlsx: Path,
    report: Path,
    *,
    sheet: str,
    out: Path | None,
) -> tuple[Path, int, int]:
    answers = _parse_report(report)
    if not answers:
        raise SystemExit(f"No answers parsed from {report}")

    wb = _load_workbook(xlsx)
    if sheet not in wb.sheetnames:
        raise SystemExit(f"Sheet not found: {sheet} (have {wb.sheetnames})")

    ws = wb[sheet]
    filled = 0
    missing: list[str] = []
    for row in range(_DATA_START, ws.max_row + 1):
        std_q = ws.cell(row=row, column=2).value
        if not std_q or not str(std_q).strip():
            continue
        key = str(std_q).strip()
        entry = answers.get(key)
        if not entry or not entry.get("answer"):
            missing.append(key)
            continue
        ws.cell(row=row, column=_COL_SYSTEM_ANSWER, value=entry["answer"])
        filled += 1

    if out is None:
        stamp = datetime.now().strftime("%Y%m%d")
        out = xlsx.with_name(f"{xlsx.stem}_RAG填答_{stamp}{xlsx.suffix}")
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    wb.close()
    return out, filled, len(answers) - filled


def main() -> int:
    parser = argparse.ArgumentParser(description="Fill RAG answers into voice script Excel")
    parser.add_argument("--xlsx", type=Path, default=_DEFAULT_XLSX)
    parser.add_argument("--report", type=Path, default=_DEFAULT_REPORT)
    parser.add_argument("--sheet", default=_SHEET)
    parser.add_argument("--out", type=Path, default=None, help="Output xlsx (default: *_RAG填答_*.xlsx)")
    args = parser.parse_args()

    if not args.xlsx.is_file():
        print(f"Excel not found: {args.xlsx}", file=sys.stderr)
        return 1
    if not args.report.is_file():
        print(f"Report not found: {args.report}", file=sys.stderr)
        return 1

    out, filled, skipped = fill_xlsx(
        args.xlsx, args.report, sheet=args.sheet, out=args.out
    )
    print(f"Filled {filled} rows → {out}")
    if skipped:
        print(f"Report entries not matched in sheet: {skipped}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
