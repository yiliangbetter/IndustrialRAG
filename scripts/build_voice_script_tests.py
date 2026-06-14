#!/usr/bin/env python3
"""Parse AI智能语音话术 Excel → test JSON + docs/AI智能客服话术测试参考答案.md."""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

_ROOT = Path(__file__).resolve().parent.parent
_DEFAULT_XLSX = Path(r"d:\dev\docs\AI智能语音话术 解析.xlsx")
_OUT_JSON = _ROOT / "data" / "voice_script_tests.json"
_OUT_MD = _ROOT / "docs" / "AI智能客服话术测试参考答案.md"

_COL = {
    "nav": 0,
    "standard_question": 1,
    "category": 2,
    "coverage": 3,
    "similar": 4,
    "standard_answer": 5,
    "system_fallback": 6,
    "sms": 7,
}

_SKIP_NAV = re.compile(r"暂不录入|暂不收录")
_TEMPLATE_Q = re.compile(r"XX|…|\.{2,}")
_TRANSFER_MARK = re.compile(r"不覆盖|转人工")
_OPENING_MARK = re.compile(r"NA|通用问题")


_DEFAULT_SHEETS = ("技术问题",)


def _load_rows(path: Path, *, sheets: tuple[str, ...] = _DEFAULT_SHEETS) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        import subprocess

        subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
        from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    rows: list[dict] = []
    next_id = 1
    sheet_order = [n for n in sheets if n in wb.sheetnames]
    if not sheet_order:
        sheet_order = list(wb.sheetnames)
    seen_keys: set[str] = set()
    for sheet_name in sheet_order:
        ws = wb[sheet_name]
        all_rows = list(ws.iter_rows(values_only=True))
        if len(all_rows) < 3:
            continue
        group = "商务" if "商务" in sheet_name else "技术" if "技术" in sheet_name else "综合"
        for raw in all_rows[2:]:
            cells = [str(c).strip() if c is not None else "" for c in raw]
            if len(cells) < 6:
                continue
            std_q = cells[_COL["standard_question"]]
            std_a = cells[_COL["standard_answer"]]
            if not std_q or std_q.startswith("【批量导入"):
                continue
            if not std_a and not cells[_COL["similar"]]:
                continue
            cat = cells[_COL["category"]]
            item_group = group
            if sheet_name == "总表":
                if cat in ("封边", "电脑锯", "排钻", "数控", "所有"):
                    item_group = "技术"
                elif std_q in (
                    "修边不好？",
                    "漏胶",
                    "齐头烂板",
                    "仿形效果不好",
                    "温控器坏了",
                    "气压报警？",
                ) or cat in ("封边", "电脑锯", "排钻", "数控"):
                    item_group = "技术"
                elif not cells[_COL["nav"]] and cat:
                    item_group = "技术"
                else:
                    item_group = "商务"
            dedupe_key = f"{item_group}:{std_q}"
            if dedupe_key in seen_keys:
                continue
            seen_keys.add(dedupe_key)
            nav = cells[_COL["nav"]]
            skip = bool(_SKIP_NAV.search(nav))
            coverage = cells[_COL["coverage"]]
            if _OPENING_MARK.search(coverage):
                expect = "script"
            elif _TRANSFER_MARK.search(coverage):
                expect = "transfer"
            elif item_group == "技术":
                expect = "technical"
            else:
                expect = "script"
            utterances = _split_utterances(cells[_COL["similar"]], std_q)
            rows.append(
                {
                    "id": next_id,
                    "group": item_group,
                    "nav": nav,
                    "category": cells[_COL["category"]],
                    "standard_question": std_q,
                    "utterances": utterances,
                    "standard_answer": std_a,
                    "system_fallback": cells[_COL["system_fallback"]],
                    "sms_template": cells[_COL["sms"]],
                    "coverage": coverage,
                    "expect": expect,
                    "skip": skip,
                    "grade": _grade_spec(std_a, expect, cells[_COL["system_fallback"]]),
                }
            )
            next_id += 1
    wb.close()
    return rows


def _split_utterances(similar: str, standard_question: str) -> list[str]:
    parts: list[str] = []
    seen: set[str] = set()

    def add(text: str) -> None:
        text = text.strip()
        if not text or text in ("-", "直接转人工"):
            return
        if _TEMPLATE_Q.search(text) and text != standard_question:
            return
        key = re.sub(r"\s+", "", text)
        if key in seen:
            return
        seen.add(key)
        parts.append(text)

    for chunk in re.split(r"##|[\n\r/]+", similar or ""):
        add(chunk)
    if not _TEMPLATE_Q.search(standard_question):
        add(standard_question)
    if not parts and standard_question:
        if _TEMPLATE_Q.search(standard_question):
            samples = {
                "我是XX省XX市 … … … …": "我是广东省广州市",
                "我是XX省XX市,如：广西、云南（精确到省/到市）": "我是云南省昆明市",
                "东莞市": "东莞市",
            }
            parts.append(samples.get(standard_question, standard_question))
        else:
            parts.append(standard_question)
    return parts[:6]


def _grade_spec(answer: str, expect: str, fallback: str) -> dict:
    text = (answer or "").strip()
    numbered = re.findall(r"(?:\d+[\.、．]|\d+\))\s*([^\d\n；;]+)", text)
    phrases: list[str] = []
    for item in numbered:
        item = item.strip().rstrip("。")
        if len(item) >= 4:
            phrases.append(item[:40])
    if not phrases:
        for seg in re.split(r"[。；;\n]", text):
            seg = re.sub(r"^\d+[\.、．]\s*", "", seg.strip())
            if len(seg) >= 6:
                phrases.append(seg[:40])
    phrases = phrases[:6]
    spec: dict = {
        "text_any": [],
        "text_all": [],
        "accept_fallback": expect == "transfer",
    }
    if expect == "technical" and phrases:
        spec["text_any"] = phrases[:4]
        if len(phrases) >= 2:
            spec["text_all"] = [phrases[0]]
    elif phrases:
        spec["text_any"] = phrases[:3]
    else:
        for token in re.findall(r"[\u4e00-\u9fff]{4,}", text):
            if token not in spec["text_any"]:
                spec["text_any"].append(token)
            if len(spec["text_any"]) >= 3:
                break
    if fallback:
        spec["fallback_any"] = ["人工客服", "转人工", "专线"]
    return spec


def _write_md(cases: list[dict], path: Path) -> None:
    lines = [
        "# AI 智能客服话术测试参考答案",
        "",
        "> 来源：`d:\\dev\\docs\\AI智能语音话术 解析.xlsx`",
        "> 生成：`python scripts/build_voice_script_tests.py`",
        "",
        "## 说明",
        "",
        "- **script**：应答须命中标准话术要点（开场白、固定回复）",
        "- **transfer**：商务流程类；可接受标准话术 **或** 转人工兜底（含「人工客服」「专线」）",
        "- **technical**：技术故障类；应答须覆盖标准答案中的排查要点",
        "- **skip**：导航标注「暂不录入/暂不收录」的条目，批测默认跳过",
        "- 本批测仅收录 Excel **技术问题** 工作表（不含商务问题）",
        "",
        f"- 用例总数：**{len(cases)}**（可测 {sum(1 for c in cases if not c['skip'])}）",
        "",
    ]
    current_group = ""
    for case in cases:
        if case["group"] != current_group:
            current_group = case["group"]
            lines.extend(["---", "", f"## {current_group}话术", ""])
        flag = "（跳过）" if case["skip"] else ""
        lines.append(f"### VS{case['id']:02d} {case['standard_question']}{flag}")
        lines.append("")
        if case["category"]:
            lines.append(f"- **分类**：{case['category']}")
        lines.append(f"- **期望类型**：`{case['expect']}`")
        if case["nav"]:
            lines.append(f"- **导航**：{case['nav']}")
        if case["coverage"]:
            lines.append(f"- **文档覆盖**：{case['coverage']}")
        lines.append(f"- **测试问法**：{'；'.join(case['utterances'][:3])}")
        lines.append("")
        lines.append("**标准答案**")
        lines.append("")
        lines.append(case["standard_answer"] or "（无）")
        lines.append("")
        if case.get("system_fallback"):
            lines.append(f"- **兜底话术**：{case['system_fallback']}")
        if case.get("sms_template"):
            lines.append(f"- **短信模板**：{case['sms_template']}")
        grade = case.get("grade") or {}
        if grade.get("text_any"):
            lines.append(f"- **批测要点（any）**：{', '.join(grade['text_any'][:4])}")
        lines.append("")

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description="Build voice script test cases from Excel")
    parser.add_argument("--xlsx", type=Path, default=_DEFAULT_XLSX)
    parser.add_argument("--out-json", type=Path, default=_OUT_JSON)
    parser.add_argument("--out-md", type=Path, default=_OUT_MD)
    parser.add_argument(
        "--all-sheets",
        action="store_true",
        help="Include 总表/商务问题/技术问题 (default: 技术问题 only)",
    )
    args = parser.parse_args()
    if not args.xlsx.is_file():
        print(f"Excel not found: {args.xlsx}", file=sys.stderr)
        return 1
    sheets = ("总表", "商务问题", "技术问题") if args.all_sheets else _DEFAULT_SHEETS
    cases = _load_rows(args.xlsx, sheets=sheets)
    payload = {
        "schema": "voice_script_tests_v1",
        "source_xlsx": str(args.xlsx),
        "count": len(cases),
        "cases": cases,
    }
    args.out_json.parent.mkdir(parents=True, exist_ok=True)
    args.out_json.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    _write_md(cases, args.out_md)
    active = sum(1 for c in cases if not c["skip"])
    print(f"Wrote {len(cases)} cases ({active} active) → {args.out_json}")
    print(f"Wrote markdown → {args.out_md}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
